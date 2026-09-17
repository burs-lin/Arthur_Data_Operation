"""单元测试 + 端到端:边框与样式范围(有效范围)。

覆盖以下场景:
  - 基础:多行表头,有效范围 = (1, 5, 1, 3)
  - 前后空行空列:有效范围 = (2, 4, 2, 4)
  - xlsm 被拒绝
"""
import sys
sys.path.insert(0, r"D:/Trae_Work/Data_Operation/.agents/skills/excel-data-cleaner/scripts")
from pathlib import Path
from openpyxl import Workbook, load_workbook
import clean_excel as cec


TMP = Path(r"D:/Trae_Work/Data_Operation/.agents/skills/excel-data-cleaner/tests/_tmp")
TMP.mkdir(exist_ok=True)


def make_xlsx(path: Path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


def get_border_map(ws):
    """返回 {(r,c): bool} 表示该单元格是否有边框。"""
    result = {}
    for row in ws.iter_rows():
        for cell in row:
            b = cell.border
            has = bool(b and b.left.style and b.right.style and b.top.style and b.bottom.style)
            result[(cell.row, cell.column)] = has
    return result


def get_merged_cells(ws):
    """返回被合并的从属单元格集合 {(r,c)}。"""
    merged = set()
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) != (mr.min_row, mr.min_col):
                    merged.add((r, c))
    return merged


def assert_border_inside_rect(rect, bmap, merged):
    fr, lr, fc, lc = rect
    for r in range(fr, lr + 1):
        for c in range(fc, lc + 1):
            if (r, c) in merged:
                continue
            assert bmap.get((r, c)), f"({r},{c}) should have border, rect={rect}"


def assert_border_outside_rect(rect, bmap):
    fr, lr, fc, lc = rect
    for (r, c), has in bmap.items():
        in_rect = (fr <= r <= lr) and (fc <= c <= lc)
        if not in_rect:
            assert not has, f"({r},{c}) should NOT have border (outside rect {rect})"


def test_e2e_basic():
    """基础:多行表头,data_rect = (1, 5, 1, 3)"""
    src = TMP / "t1.xlsx"
    make_xlsx(src, [
        ["销售报表", None, None],
        ["区域", "门店", "GMV"],
        ["华北", "北京店", "合计"],
        ["华南", "深圳店", 500],
        ["华东", "上海店", 800],
    ])
    out = TMP / "t1_out.xlsx"
    cec.clean_file(str(src), output_path=str(out))
    wb = load_workbook(out)
    ws = wb["Sheet1"]
    rect = cec._find_effective_range(ws)
    bmap = get_border_map(ws)
    merged = get_merged_cells(ws)
    assert rect == (1, 5, 1, 3), f"expected (1,5,1,3), got {rect}"
    assert_border_inside_rect(rect, bmap, merged)
    assert_border_outside_rect(rect, bmap)
    print(f"[OK] test_e2e_basic rect={rect}")


def test_e2e_empty_margin():
    """前后空行空列:data_rect = (2, 4, 2, 4)"""
    src = TMP / "t2.xlsx"
    make_xlsx(src, [
        [None, None, None, None],
        [None, "姓名", "分数", "日期"],
        [None, "张三", 90, "2024-01-01"],
        [None, "李四", 85, "2024-01-02"],
        [None, None, None, None],
    ])
    out = TMP / "t2_out.xlsx"
    cec.clean_file(str(src), output_path=str(out))
    wb = load_workbook(out)
    ws = wb["Sheet1"]
    rect = cec._find_effective_range(ws)
    bmap = get_border_map(ws)
    merged = get_merged_cells(ws)
    assert rect == (2, 4, 2, 4), f"expected (2,4,2,4), got {rect}"
    assert_border_inside_rect(rect, bmap, merged)
    assert_border_outside_rect(rect, bmap)
    print(f"[OK] test_e2e_empty_margin rect={rect}")


def test_xlsm_rejected():
    """.xlsm 必须被拒绝。"""
    src = TMP / "t3.xlsm"
    make_xlsx(src, [["a", "b"], [1, 2]])
    try:
        cec.clean_file(str(src), output_path=str(TMP / "t3_out.xlsx"))
        raise AssertionError("xlsm 应该被拒绝,但通过了")
    except ValueError as e:
        assert ".xlsm" in str(e)
        print(f"[OK] test_xlsm_rejected: {e}")


if __name__ == "__main__":
    test_e2e_basic()
    test_e2e_empty_margin()
    test_xlsm_rejected()
    print("\nALL DATA RECT TESTS PASSED")